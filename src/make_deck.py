"""Build outputs/Milestone1_Sumukh_S24CSEU0570.pptx (10 slides) from the generated JSON/PNG outputs.

Numbers come from outputs/eda_stats.json and outputs/baseline_results.json; examples come from the pilot CSV by id.
Run: python src/make_deck.py   (after eda.py and baseline.py)
"""
import json
import math
import os

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

OUT = "outputs/"
INK, MUTED = RGBColor(0x1B, 0x26, 0x3B), RGBColor(0x5C, 0x67, 0x7D)
NAVY, TEAL, CORAL, GREY = RGBColor(0x1F, 0x2A, 0x44), RGBColor(0x2A, 0x9D, 0x8F), RGBColor(0xE7, 0x6F, 0x51), RGBColor(0x8D, 0x99, 0xAE)
PANEL, WHITE, AMBER = RGBColor(0xF1, 0xF4, 0xF8), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE9, 0xA8, 0x20)
LABEL_COLOR = {"Favour": TEAL, "Against": CORAL, "None": GREY}
WARN = []

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def _est_lines(text, width_in, size):
    wide = any(0x0900 <= ord(c) <= 0x09FF for c in text)  # Indic glyphs are wider than Latin
    cpl = max(1, int(width_in * 72 / (size * (0.62 if wide else 0.5))))
    return max(1, sum(math.ceil(max(1, len(p)) / cpl) for p in text.split("\n")))


def _style(run, size, bold, color, italic=False):
    f = run.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = color
    f.name = "Calibri"
    rpr = run._r.get_or_add_rPr()
    cs = etree.SubElement(rpr, qn("a:cs"))
    cs.set("typeface", "Nirmala UI")  # complex-script font so Devanagari/Bengali render on Windows lab PCs


def text(slide, x, y, w, h, paras, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         fill=None, line=None, pad=0.1, space=4, shape=MSO_SHAPE.RECTANGLE):
    """paras: list of str | list[(text, {size,bold,color,italic})] runs. Adds a shape with a text frame."""
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.5)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    need = 0.0
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space)
        runs = [(p, {})] if isinstance(p, str) else p
        plain, psize = "", size
        for t, o in runs:
            r = para.add_run()
            r.text = t
            psize = o.get("size", size)
            _style(r, psize, o.get("bold", bold), o.get("color", color), o.get("italic", False))
            plain += t
        need += _est_lines(plain, w - 2 * pad, psize) * psize * 1.22 / 72 + space / 72
    if need > h:
        WARN.append(f"slide {len(prs.slides)}: text may overflow ({need:.2f}in > {h:.2f}in): {paras[0] if isinstance(paras[0], str) else paras[0][0][0]!r:.50}")
    return sh


def image(slide, path, x, y, w=None, h=None):
    iw, ih = Image.open(path).size
    if w is not None:
        h = w * ih / iw
    else:
        w = h * iw / ih
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def table(slide, x, y, w, col_w, rows, size=15, header_fill=NAVY, row_h=0.5):
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = shp.table
    tbl.horz_banding = False
    total = sum(col_w)
    heights = 0.0
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(w * cw / total)
    for i, row in enumerate(rows):
        worst = 1
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else PANEL)
            tf = c.text_frame
            tf.word_wrap = True
            r = tf.paragraphs[0].add_run()
            r.text = str(val)
            _style(r, size, i == 0 or j == 0, WHITE if i == 0 else INK)
            worst = max(worst, _est_lines(str(val), w * col_w[j] / total - 0.16, size))
        rh = worst * size * 1.22 / 72 + 0.12
        tbl.rows[i].height = Inches(max(rh, 0.38))
        heights += max(rh, 0.38)
    if y + heights > 6.85:
        WARN.append(f"slide {len(prs.slides)}: table may run past the footer (ends {y + heights:.2f}in)")
    return heights


def new_slide(num, tag, title, notes=""):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.55, 0.28, 9.5, 0.4, [tag.upper()], size=14, color=TEAL, bold=True, pad=0)
    text(s, 0.55, 0.62, 12.2, 0.95, [title], size=30, color=NAVY, bold=True, pad=0, anchor=MSO_ANCHOR.TOP)
    text(s, 0.55, 7.05, 10, 0.3, ["Sumukh Gupta, Neya Tripathi, Aditya Goel, Kushika Mishra  |  CSET346 NLP  |  Milestone 1"], size=11, color=MUTED, pad=0)
    text(s, 11.6, 7.05, 1.2, 0.3, [f"{num} / 10"], size=11, color=MUTED, pad=0, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def pct(x):
    return f"{x:.1f}%"


# ------------------------------------------------------------------ data
E = json.load(open(OUT + "eda_stats.json"))
B = json.load(open(OUT + "baseline_results.json"))
P = pd.read_csv("data/pilot_indic_stance.csv", dtype=str, keep_default_na=False).set_index("id")
G = {(r["features"], r["model"]): r for r in B["A_english_cv"]}
best = G[("word+char", "LogReg")]
floor = G[("-", "Dummy(majority)")]
abl = B["A_class_weight_ablation"]
D = B["D_transfer_to_pilot"]
eng, pil, X = E["english_merged"], E["pilot"], E["extra"]
ex = lambda i: P.loc[i, "comment"]  # noqa: E731

# ------------------------------------------------------------------ 1 title
s = prs.slides.add_slide(BLANK)
text(s, 0, 0, 13.333, 7.5, [""], fill=NAVY)
text(s, 0.8, 0.7, 11, 0.5, ["CSET346  NATURAL LANGUAGE PROCESSING  |  MILESTONE 1"], size=16, color=TEAL, bold=True, pad=0)
text(s, 0.8, 1.5, 11.8, 2.5, ["Detecting Stance in Media on Global Warming for Indic Languages"], size=46, color=WHITE, bold=True, pad=0)
text(s, 0.8, 3.9, 11.5, 0.9, ["Favour / Against / None toward one claim, in Hindi and Bengali comments"], size=24, color=RGBColor(0xC9, 0xD3, 0xE3), pad=0)
text(s, 0.8, 4.85, 3.6, 2.2, ["Sumukh Gupta  |  S24CSEU0570", "Neya Tripathi  |  S24CSEU0609", "Aditya Goel  |  S24CSEU0614", "Kushika Mishra  |  S24CSEU0652"], size=17, color=WHITE, pad=0)
text(s, 4.7, 4.85, 3.7, 2.2, ["B.Tech CSE, Bennett University", "Instructor: Dr. Supriya Chanda"], size=17, color=RGBColor(0xC9, 0xD3, 0xE3), pad=0)
text(s, 8.6, 5.2, 4.2, 1.4, ["Derived from the AISoMe 2026 shared task at FIRE 2026"], size=16, color=RGBColor(0xC9, 0xD3, 0xE3), pad=0)
s.notes_slide.notes_text_frame.text = "Introduce: project, claim, that the task comes from AISoMe 2026 (English train, Hindi/Bengali test)."

# ------------------------------------------------------------------ 2 problem identification
s = new_slide(2, "1 | Problem identification", "One claim, three labels, Hindi and Bengali comments",
              "The claim never appears in the comments. Examples are from our team-authored pilot corpus, not scraped data.")
text(s, 0.55, 1.5, 12.2, 0.8, ["“Climate change and global warming is a serious concern.”"], size=26, color=NAVY, bold=True,
     fill=PANEL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
cols = [("Favour", "Agrees it is a serious concern", "P003", "P078"),
        ("Against", "Disagrees, minimises, denies", "P027", "P102"),
        ("None", "Neutral, off-topic or unclear", "P057", "P132")]
for k, (lab, desc, hi, bn) in enumerate(cols):
    x = 0.55 + k * 4.12
    text(s, x, 2.55, 3.95, 0.6, [lab], size=26, bold=True, color=WHITE, fill=LABEL_COLOR[lab], anchor=MSO_ANCHOR.MIDDLE,
         align=PP_ALIGN.CENTER)
    text(s, x, 3.2, 3.95, 0.5, [desc], size=16, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, x, 3.75, 3.95, 1.45, [[("Hindi  ", {"bold": True, "size": 14, "color": MUTED})], ex(hi)], size=19, fill=PANEL)
    text(s, x, 5.3, 3.95, 1.45, [[("Bengali  ", {"bold": True, "size": 14, "color": MUTED})], ex(bn)], size=19, fill=PANEL)
text(s, 0.55, 6.72, 12.2, 0.3, ["Examples: team-authored pilot corpus (synthetic, development only). Metric: macro-F1."], size=12, color=MUTED, pad=0)

# ------------------------------------------------------------------ 3 problem analysis
s = new_slide(3, "1 | Problem analysis", "Why this is not simple text classification",
              "Four arguments: external claim, shared vocabulary, implicit stance, heterogeneous None. Stats are measured.")
cards = [
    ("1  The claim is not in the text", "Against, because the claim is about warming:", ex("P046") + "  (Against)", None),
    ("2  Same words, opposite stance", f"{X['favour_against_vocab_overlap_pct']}% of Against vocabulary also appears in Favour (English set).",
     ex("P010") + "  (Favour)\n" + ex("P027") + "  (Against)", None),
    ("3  Stance is often implicit", "World knowledge needed; no climate word:", ex("P019") + "  (Favour)", None),
    ("4  None is a junk drawer", f"Neutral + off-topic + ambiguous. {X['raw_ties']} of {X['raw_total']:,} annotated sentences "
     f"({pct(X['raw_ties_pct'])}) had no annotator majority.", ex("P051") + "  (None)", None),
]
for k, (h, body, exm, _) in enumerate(cards):
    x, y = 0.55 + (k % 2) * 6.15, 1.55 + (k // 2) * 2.7
    text(s, x, y, 6.0, 0.5, [h], size=21, bold=True, color=NAVY, pad=0.12)
    text(s, x, y + 0.5, 6.0, 0.7, [body], size=16, color=MUTED, pad=0.12)
    text(s, x, y + 1.2, 6.0, 1.35, [exm], size=15, fill=PANEL, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.8, 12.2, 0.25, ["Plus: English training data, Hindi/Bengali test data. Examples: pilot corpus (synthetic)."], size=11, color=MUTED, pad=0)

# ------------------------------------------------------------------ 4 existing systems
s = new_slide(4, "2 | Review of existing systems", "What exists, what it reports, why it falls short",
              "Numbers are as reported by each source on its own data; not comparable with each other or with ours.")
rows = [["Approach", "Representative work", "Reported result", "Limitation for us"],
        ["Tweet stance, 3 labels", "SemEval-2016 Task 6 (Mohammad)", "see paper", "English; per-target"],
        ["News-sentence stance", "GWSD (Luo et al. 2020)", "see paper", "English news, not comments"],
        ["Climate tweets, 6 layers", "ClimaConvo (Shiwakoti 2024)", "15,309 rows", "IDs only: no text"],
        ["Multimodal video", "MultiClimate (Wang 2024)", "acc 0.747, F1 0.749", "English transcripts"],
        ["Bangla news stance", "Dhoroni (Wasi 2024)", "acc 0.635, F1 0.308", "News; minority classes weak"],
        ["Code-mixed Hindi", "Swami et al. 2018", "58.7% accuracy", "Target: demonetisation"],
        ["Same task, translate-train", "Squad4, AISoMe 2026", "F1 0.42-0.53 (val 0.65-0.68)", "Large val-to-test gap"]]
table(s, 0.55, 1.55, 12.2, [2.5, 3.0, 2.8, 2.7], rows, size=16, row_h=0.6)
text(s, 0.55, 6.62, 12.2, 0.4, ["Squad4: a competing team; only their reported scores are cited. Ours (measured): TF-IDF macro-F1 "
                                f"{best['macro_f1']:.3f} on English CV."], size=13, color=MUTED, pad=0)

# ------------------------------------------------------------------ 5 research gap
s = new_slide(5, "2 | Research gap", "No resource matches target, language and domain",
              "Headline measured result: English-trained TF-IDF on the Indic pilot gives macro-F1 0.341 and predicts None for 131/150. "
              "Gap table: data/DATASET_INVENTORY.md. Pilot is synthetic, development only.")
text(s, 0.55, 1.55, 3.6, 0.9, ["English labelled data"], size=22, bold=True, color=WHITE, fill=TEAL, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 2.5, 3.6, 0.6, ["GWSD news + SemEval tweets"], size=16, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 1.75, 3.05, 1.2, 0.75, ["↓"], size=34, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.55, 3.85, 3.6, 0.9, ["Hindi / Bengali comments"], size=22, bold=True, color=WHITE, fill=CORAL, align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 4.8, 3.6, 0.9, ["hidden test: code-mixing, spelling variation, negation, implicit stance"], size=14, color=MUTED,
     align=PP_ALIGN.CENTER)
rows = [["Resource", "Our claim", "Hindi / Bengali", "Comments"],
        ["SemEval-2016 climate", "Yes", "No", "Yes"],
        ["GWSD", "Yes", "No", "No (news)"],
        ["ClimaConvo", "Close", "No", "Yes, but no text"],
        ["MultiClimate", "Close", "No", "No (transcripts)"],
        ["Dhoroni", "Close", "Bengali", "No (news)"],
        ["Swami 2018", "No (demonetisation)", "Romanised Hindi", "Yes"],
        ["AISoMe test (hidden)", "Yes", "Yes", "Yes"]]
table(s, 4.6, 1.55, 8.15, [2.6, 2.2, 2.0, 2.2], rows, size=15, row_h=0.5)
text(s, 0.55, 5.85, 12.2, 1.1, [[("Measured, not assumed: ", {"bold": True, "color": CORAL}),
     (f"English-trained TF-IDF on the Indic pilot reaches macro-F1 {D['macro_f1']:.3f} (floor 0.167) and predicts None for "
      f"{D['predicted_class_counts']['None']} of 150: cross-lingual transfer collapses to majority-class prediction, so "
      f"in-language data and Indic-specific models are needed. {X['favour_against_vocab_overlap_pct']}% of Against vocabulary "
      "also appears in Favour: not simple text classification. ", {}),
     ("(Pilot: synthetic, development only.)", {"italic": True, "color": MUTED})]],
     size=15, fill=RGBColor(0xFD, 0xEC, 0xE8), line=CORAL, pad=0.15, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------------ 6 dataset exploration
s = new_slide(6, "3 | Dataset exploration and statistics", "What data we actually have",
              "Usable dataset: PARTIALLY. English target-correct data yes; Hindi/Bengali comments only the synthetic pilot.")
image(s, OUT + "fig_class_distribution.png", 0.45, 1.55, w=7.0)
text(s, 7.7, 1.55, 5.1, 0.55, ["Usable dataset found: PARTIALLY"], size=22, bold=True, color=WHITE, fill=AMBER,
     anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
text(s, 7.7, 2.25, 5.1, 4.6, [
    [("English merged: ", {"bold": True}), (f"{eng['n']:,} labelled (GWSD news + SemEval tweets)", {})],
    [("Against is rare: ", {"bold": True}), (f"{eng['class_pct']['Against']}% (SemEval: 26 of 564)", {})],
    [("Hapax rate: ", {"bold": True}), (f"{eng['hapax_rate_pct']}% English, {pil['hapax_rate_pct']}% pilot", {})],
    [("Majority floor: ", {"bold": True}), (f"macro-F1 {eng['majority_floor_macro_f1']:.3f}", {})],
    [("No annotator majority: ", {"bold": True}), (f"{X['raw_ties']} of {X['raw_total']:,} GWSD ({pct(X['raw_ties_pct'])}); "
                                                f"{X['ties']} remain after dedup, dropped", {})],
    [("Median length: ", {"bold": True}), (f"{eng['words']['median']:.0f} words (pilot {pil['words']['median']:.0f})", {})],
    [("Hindi/Bengali: ", {"bold": True, "color": CORAL}), ("150 team-authored synthetic comments, development only", {})],
], size=18, space=7)
text(s, 0.45, 5.3, 7.0, 1.4, ["Not obtained: official AISoMe data; ClimaConvo text (IDs only). Dhoroni (Bengali news) and Swami "
                              "(demonetisation) kept out of the training pool."], size=15, color=MUTED, pad=0.1)

# ------------------------------------------------------------------ 7 objectives and methodology
s = new_slide(7, "3 | Objectives and methodology", "Pipeline: clean text to cross-lingual evaluation",
              "Green = implemented and run today. Amber = written, not yet run. Grey = planned for Milestone 2.")
text(s, 0.55, 1.5, 12.2, 1.2, [
    "1  Understand and measure the task and data honestly",
    "2  Build a reproducible baseline with a floor to beat",
    "3  Design a target-aware multilingual model and a cross-lingual strategy"], size=19, color=INK, space=3)
steps = [("Preprocess", "NFC, noise strip, script ID", TEAL, "done"),
         ("Baselines", "TF-IDF word + char, majority floor", TEAL, "done"),
         ("Target-aware encoder", "MuRIL on claim [SEP] comment", AMBER, "written"),
         ("Cross-lingual", "zero-shot vs translate-train", GREY, "plan"),
         ("Evaluate", "macro-F1, per phenomenon", GREY, "plan")]
for k, (h, sub, col, st) in enumerate(steps):
    x = 0.55 + k * 2.5
    text(s, x, 3.2, 2.2, 0.95, [h], size=20, bold=True, color=WHITE, fill=col, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, 4.25, 2.2, 1.0, [sub], size=16, color=INK, align=PP_ALIGN.CENTER)
    text(s, x, 5.25, 2.2, 0.4, [st], size=14, color=col, bold=True, align=PP_ALIGN.CENTER)
    if k < 4:
        text(s, x + 2.17, 3.35, 0.4, 0.6, ["→"], size=28, color=NAVY, bold=True, pad=0, align=PP_ALIGN.CENTER)
text(s, 0.55, 5.9, 12.2, 0.9, ["Models are compared against the floor; the pilot is for development and error analysis only."], size=16, color=MUTED, pad=0)

# ------------------------------------------------------------------ 8 algorithms
s = new_slide(8, "4 | Relevance of the techniques", "Why each choice, with evidence",
              "Char n-grams: hapax rate. MuRIL: IndicXNLI. Pair input: stance is relative to a claim. Macro-F1 + weights: rare Against.")
cards = [("Char n-grams (2-5)", f"Hapax rate {eng['hapax_rate_pct']}% English, {pil['hapax_rate_pct']}% pilot: most words seen once. "
          "Sub-strings are shared across spellings and Romanised text."),
         ("MuRIL encoder", "IndicXNLI: MuRIL beats XLM-R, mBERT and IndicBERT on Indic languages; English + Indic training is best. "
          "Pretrained with transliterated text."),
         ("claim [SEP] comment", "Stance is defined relative to the claim, so the model must see both, not just the topic words."),
         ("Macro-F1 + class weights", f"Floor {floor['macro_f1']:.2f}. Weights: macro-F1 {abl['none']['macro_f1']:.2f} to "
          f"{abl['balanced']['macro_f1']:.2f}, Against recall {abl['none']['against_recall']:.2f} to {abl['balanced']['against_recall']:.2f}.")]
for k, (h, body) in enumerate(cards):
    x, y = 0.55 + (k % 2) * 6.15, 1.55 + (k // 2) * 2.7
    text(s, x, y, 6.0, 0.6, [h], size=24, bold=True, color=WHITE, fill=[TEAL, NAVY, NAVY, TEAL][k], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, y + 0.65, 6.0, 1.85, [body], size=22, fill=PANEL, pad=0.15)

# ------------------------------------------------------------------ 9 results
s = new_slide(9, "Initial implementation and preliminary results", "What runs today",
              "Headline: English-to-Indic zero-shot collapses to majority-class prediction (pilot: synthetic, dev only). "
              "English CV numbers are optimistic (no document ids, one hashtag cluster).")
image(s, OUT + "fig_baseline_comparison.png", 0.45, 1.45, w=5.9)
image(s, OUT + "fig_confusion_english_cv.png", 6.5, 1.45, w=3.2)
le = B["A_best"]["largest_error"]
text(s, 9.9, 1.6, 3.0, 2.4, ["Best (English CV):", "word+char LogReg", f"macro-F1 {best['macro_f1']:.3f}", f"floor {floor['macro_f1']:.3f}"],
     size=17, color=NAVY, bold=True, pad=0.05)
text(s, 0.55, 4.8, 12.2, 1.3, [[("Headline finding: ", {"bold": True, "color": CORAL}),
     (f"English to Indic zero-shot collapses: macro-F1 {D['macro_f1']:.3f} vs floor 0.167, predicting None for "
      f"{D['predicted_class_counts']['None']} of 150 comments. With {X['favour_against_vocab_overlap_pct']}% of Against vocabulary "
      "also in Favour, this is not simple text classification. ", {}),
     ("(Pilot: synthetic, dev only; not comparable to official results.)", {"italic": True, "color": MUTED})]],
     size=17, fill=RGBColor(0xFD, 0xEC, 0xE8), line=CORAL, pad=0.15, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.55, 6.2, 12.2, 0.8, [
    [("Runs today: ", {"bold": True}), ("loader + validation gate, preprocessing, EDA, TF-IDF baselines, self-checks.", {})],
    [("Error analysis: ", {"bold": True}), (f"largest confusion is true {le['true']} predicted {le['pred']} ({le['n']}): factual "
                                           "sentences with no opinion marker. Against recall "
                                           f"{B['A_best']['against_recall']:.2f}.", {})]], size=14, space=2)

# ------------------------------------------------------------------ 10 feasibility
s = new_slide(10, "2 | Feasibility, risks and next steps", "Feasible, with known risks",
              "Biggest risk: validation-to-test gap reported by a peer team (0.65-0.68 to 0.42-0.53).")
cols = [("Compute", ["Baselines run on a laptop CPU in minutes (done)",
                     "Fine-tuning a base-size encoder on short texts: Colab GPU planned",
                     "Runtime: not yet measured"]),
        ("Risks and mitigation", ["Official data not in hand: obtain from the course email; pilot for development only",
                                       "Against is rare: class weights, macro-F1",
                                       "Validation-to-test gap: per-phenomenon evaluation"]),
        ("Milestone 2", ["Run MuRIL on claim [SEP] comment (XLM-R as comparison)",
                         "Zero-shot vs translate-train vs both",
                         "Zero-shot LLM baseline",
                         "Swap in official data"])]
for k, (h, items) in enumerate(cols):
    x = 0.55 + k * 4.12
    text(s, x, 1.6, 3.95, 0.6, [h], size=22, bold=True, color=WHITE, fill=[TEAL, CORAL, NAVY][k], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, 2.3, 3.95, 4.3, items, size=20, fill=PANEL, pad=0.15, space=12)

# ------------------------------------------------------------------ save
os.makedirs(OUT, exist_ok=True)
path = OUT + "Milestone1_Sumukh_S24CSEU0570.pptx"
assert len(prs.slides) <= 10
prs.save(path)
print(f"saved {path} with {len(prs.slides)} slides")
for w in WARN:
    print("WARNING", w)
print("fit warnings:", len(WARN))
